#!/usr/bin/env python3
"""Bouwt de Foert prijs-rekenhulp (Excel)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ORANGE = "FF5A1F"
DARK = "15171C"
LIGHT = "FFF1EA"
GREY = "F4F6F9"
WHITE = "FFFFFF"

bold = Font(bold=True, color=DARK)
white_bold = Font(bold=True, color=WHITE, size=12)
title_font = Font(bold=True, color=WHITE, size=16)
section_font = Font(bold=True, color=WHITE, size=11)
note_font = Font(italic=True, color="6B6B6B", size=9)

orange_fill = PatternFill("solid", fgColor=ORANGE)
dark_fill = PatternFill("solid", fgColor=DARK)
light_fill = PatternFill("solid", fgColor=LIGHT)
grey_fill = PatternFill("solid", fgColor=GREY)
input_fill = PatternFill("solid", fgColor="FFF8E6")

thin = Side(style="thin", color="DDDDDD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

EUR = '#,##0.00 "€"'
PCT = '0%'

wb = Workbook()

# ===================== SHEET 1: REKENHULP =====================
ws = wb.active
ws.title = "Rekenhulp"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 40

def cell(ref, value=None, font=None, fill=None, fmt=None, align=None, bd=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font: c.font = font
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if align: c.alignment = Alignment(horizontal=align, vertical="center")
    if bd: c.border = border
    return c

# Titel
ws.merge_cells("B2:D2")
cell("B2", "FOERT — prijs-rekenhulp", title_font, orange_fill, align="left")
ws.row_dimensions[2].height = 30
ws.merge_cells("B3:D3")
cell("B3", "Vul de gele cellen in. De rest rekent zichzelf uit.", note_font)

def section(row, text):
    ws.merge_cells(f"B{row}:D{row}")
    cell(f"B{row}", text, section_font, dark_fill, align="left")
    ws.row_dimensions[row].height = 22

def line(row, label, value, fmt=None, is_input=False, formula=False, note=None, total=False):
    cell(f"B{row}", label, bold if total else None, bd=True,
         fill=light_fill if total else None)
    c = ws[f"C{row}"]
    c.value = value
    c.border = border
    if fmt: c.number_format = fmt
    c.alignment = Alignment(horizontal="right", vertical="center")
    if is_input:
        c.fill = input_fill
        c.font = Font(bold=True, color=DARK)
    if total:
        c.fill = light_fill
        c.font = Font(bold=True, color=ORANGE, size=12)
    if note:
        cell(f"D{row}", note, note_font)

# ---- INSTELLINGEN ----
section(5, "1. INSTELLINGEN  (jouw eigen tarieven — pas aan)")
line(6,  "Stortkost per kg (recyclagepark)", 0.35, EUR, is_input=True, note="2026 ± 0,35 €/kg. Check je lokale park.")
line(7,  "Voertuigkost per km", 0.30, EUR, is_input=True, note="Brandstof + slijtage bestelwagen.")
line(8,  "Eigen uurkost per persoon", 25, EUR, is_input=True, note="Wat jullie uur jullie 'kost' (loon).")
line(9,  "Verkooptarief per m³", 35, EUR, is_input=True, note="Markt: 25–45 €/m³.")
line(10, "Verkooptarief per manuur", 50, EUR, is_input=True, note="Markt: 40–60 €/persoon/uur.")
line(11, "Gewenste marge op kostprijs", 0.40, PCT, is_input=True, note="0,40 = 40% winst op je kosten.")
line(12, "Minimumprijs per klus", 200, EUR, is_input=True, note="Onder dit bedrag niet rendabel.")
line(13, "Vuistregel kg per m³", 90, '0', is_input=True, note="Huisraad ± 80–100 kg/m³.")

# ---- JOUW KLUS ----
section(16, "2. JOUW KLUS  (vul in per opdracht)")
line(17, "Volume rommel (m³)", 30, '0.0', is_input=True, note="Schat op basis van foto's.")
line(18, "Afvalgewicht (kg) — leeg = schatting via m³", None, '0', is_input=True, note="Laat leeg om automatisch te schatten.")
line(19, "Aantal personen", 2, '0', is_input=True)
line(20, "Aantal uren op de klus", 6, '0.0', is_input=True)
line(21, "Afstand heen + terug (km)", 40, '0', is_input=True)
line(22, "Weekend / spoed toeslag", 0.0, PCT, is_input=True, note="0,20 = +20%.")
line(23, "Restwaarde inboedel (aftrek)", 0, EUR, is_input=True, note="Wat je recupereert: metaal, meubels…")
line(24, "Materiaal / diversen", 20, EUR, is_input=True, note="Zakken, klein materiaal, containerhuur…")

# ---- JOUW KOSTEN ----
section(27, "3. JOUW KOSTEN")
line(28, "Geschat afvalgewicht (kg)", '=IF(C18=\"\",C17*C13,C18)', '0', note="Automatisch of jouw invoer.")
line(29, "Stortkosten", '=C28*C6', EUR)
line(30, "Voertuig / brandstof", '=C21*C7', EUR)
line(31, "Loonkost (personen × uren × uurkost)", '=C19*C20*C8', EUR)
line(32, "Materiaal / diversen", '=C24', EUR)
line(33, "Restwaarde inboedel (−)", '=-C23', EUR)
line(34, "TOTALE KOSTPRIJS", '=SUM(C29:C33)', EUR, total=True)

# ---- RICHTPRIJS ----
section(37, "4. RICHTPRIJS  (3 methodes — kies of neem gemiddelde)")
line(38, "Methode A — per m³", '=MAX(C17*C9*(1+C22),C12)', EUR, note="Volume × tarief/m³.")
line(39, "Methode B — per uur + stortkosten", '=MAX((C19*C20*C10+C29)*(1+C22),C12)', EUR, note="Manuren × uurtarief + afvoer.")
line(40, "Methode C — kostprijs + marge", '=MAX(C34*(1+C11)*(1+C22),C12)', EUR, note="Kostprijs + jouw marge.")
line(41, "GEMIDDELDE RICHTPRIJS", '=AVERAGE(C38:C40)', EUR, total=True)

# ---- RESULTAAT ----
section(44, "5. RESULTAAT  (bij gemiddelde richtprijs)")
line(45, "Richtprijs (excl. btw)", '=C41', EUR)
line(46, "Richtprijs incl. 21% btw", '=C41*1.21', EUR, note="Indien btw-plichtig en aan particulier.")
line(47, "Jouw kostprijs", '=C34', EUR)
line(48, "WINST (€)", '=C45-C47', EUR, total=True)
line(49, "Winstmarge (%)", '=IF(C45=0,0,(C45-C47)/C45)', PCT)
line(50, "Winst per persoon", '=C48/C19', EUR)

cell("B52", "Let op: richtcijfers, geen vast tarief of fiscaal advies. Reken 1× je echte kostprijs op een proefklus.", note_font)
ws.merge_cells("B52:D52")

# ===================== SHEET 2: MARKTPRIJZEN =====================
ws2 = wb.create_sheet("Marktprijzen")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 3
for col, w in [("B", 30), ("C", 16), ("D", 22)]:
    ws2.column_dimensions[col].width = w

def head2(ws, row, cols):
    for i, t in enumerate(cols):
        c = ws.cell(row=row, column=2 + i, value=t)
        c.font = white_bold
        c.fill = orange_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border

def row2(ws, row, vals):
    for i, t in enumerate(vals):
        c = ws.cell(row=row, column=2 + i, value=t)
        c.border = border
        c.alignment = Alignment(horizontal="left", vertical="center")

ws2.merge_cells("B2:D2")
ws2["B2"] = "Marktprijzen — wat klanten betalen"
ws2["B2"].font = title_font
ws2["B2"].fill = orange_fill
ws2.row_dimensions[2].height = 28

head2(ws2, 4, ["Per eenheid", "Prijs", ""])
row2(ws2, 5, ["Per m³ rommel", "25 – 45 €", ""])
row2(ws2, 6, ["Per uur per persoon", "40 – 60 €", "+ stortkosten"])

head2(ws2, 8, ["Type woning", "Volume", "Totaalprijs"])
for r, v in enumerate([
    ["Seniorenkamer", "klein", "250 – 450 €"],
    ["Studio", "10–15 m³", "500 – 1.000 €"],
    ["Appartement", "20–35 m³", "1.000 – 2.000 €"],
    ["Gezinswoning", "40–60 m³", "1.500 – 3.000 €"],
    ["Volle woning + zolder/kelder", "80 m³+", "tot 4.000 €"],
    ["Sterk vervuild / hoarding", "—", "vanaf 1.500 €"],
], start=9):
    row2(ws2, r, v)

ws2.merge_cells("B16:D16")
ws2["B16"] = "Bron: Verhuiskampioen, Jacobs Opruimdiensten, Uw Inboedel Service e.a. (2026)"
ws2["B16"].font = note_font

# ===================== SHEET 3: TOESLAGEN & KOSTEN =====================
ws3 = wb.create_sheet("Toeslagen & kosten")
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 3
ws3.column_dimensions["B"].width = 38
ws3.column_dimensions["C"].width = 26

ws3.merge_cells("B2:C2")
ws3["B2"] = "Toeslagen, kosten & afvaltarieven"
ws3["B2"].font = title_font
ws3["B2"].fill = orange_fill
ws3.row_dimensions[2].height = 28

head2(ws3, 4, ["Verhoogt de prijs", "Bedrag"])
for r, v in enumerate([
    ["Weekend / spoedontruiming", "+20 – 40%"],
    ["Verdieping zonder lift", "extra uren / ladderlift"],
    ["Schoonmaak na ontruiming", "± 20 €/uur"],
    ["Tapijt/vloer verwijderen", "± 2,50 €/m²"],
    ["Voorrijkosten", "0,25 – 0,50 €/km"],
    ["Speciaal afval (asbest, chemisch)", "duurst, erkende verwerking"],
    ["Opslag", "± 125 €/maand (25 m³)"],
], start=5):
    row2(ws3, r, v)

head2(ws3, 14, ["Verlaagt de prijs", "Effect"])
row2(ws3, 15, ["Klant sorteert zelf voor", "−20 – 40%"])
row2(ws3, 16, ["Restwaarde inboedel", "aftrek / opkoop"])

head2(ws3, 19, ["Jouw afvalkosten 2026 (Vlaanderen)", "Tarief"])
for r, v in enumerate([
    ["Grofvuil", "± 0,35 €/kg"],
    ["Restafval", "± 0,35 €/kg (heffing verdubbelt)"],
    ["Zuiver steenpuin / harde plastics", "± 0,055 €/kg"],
], start=20):
    row2(ws3, r, v)

ws3.merge_cells("B24:C24")
ws3["B24"] = "Tarief verschilt per intercommunale (Ivarem, Imog, Limburg.net…). Check je eigen park."
ws3["B24"].font = note_font

wb.save("/home/user/Foert/Foert_prijs_rekenhulp.xlsx")
print("OK - Foert_prijs_rekenhulp.xlsx aangemaakt")
